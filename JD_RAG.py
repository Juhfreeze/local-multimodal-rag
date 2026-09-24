from __future__ import annotations

import gc
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
from functools import lru_cache
from io import BytesIO
from pathlib import Path

import ollama
from PIL import Image
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# -------------------- CONFIGURATION --------------------
BASE_DIR = Path(__file__).resolve().parent
NOTES_DIR = BASE_DIR / "my_notes"
DB_DIR = BASE_DIR / "chroma_db"
MANIFEST_FILE = BASE_DIR / "indexed_files.json"

LLM_MODEL = "qwen3:4b-instruct"  # Check exact tag with: ollama list
EMBED_MODEL = "nomic-embed-text"
VISION_MODEL = "granite3.2-vision"
COLLECTION_NAME = "my_notes"
BATCH_SIZE = 4
RETRIEVAL_K = 4
PIPELINE_VERSION = "docling-granite-v1"  # Change => rebuild index
SUPPORTED = {".txt", ".md", ".pdf", ".pptx", ".docx"}

# When LibreOffice and PyMuPDF are available, send visually meaningful slides to Granite.
# Otherwise, send embedded pictures (and extract native chart data as text).
RENDER_FULL_SLIDES = True
MAX_VISUALS_PER_FILE = 80  # Reduce for very image-heavy documents.

splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
vision_used = False


# -------------------- MANIFEST --------------------
def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_manifest() -> dict:
    if not MANIFEST_FILE.exists():
        return {"pipeline_version": PIPELINE_VERSION, "files": {}}
    with MANIFEST_FILE.open("r", encoding="utf-8") as f:
        manifest = json.load(f)
    if (not isinstance(manifest, dict)
            or manifest.get("pipeline_version") != PIPELINE_VERSION
            or not isinstance(manifest.get("files"), dict)):
        raise RuntimeError(
            "The existing manifest belongs to an older indexing pipeline. "
            "Back up and delete BOTH chroma_db and indexed_files.json, then restart."
        )
    return manifest


def save_manifest(manifest: dict) -> None:
    temp = MANIFEST_FILE.with_suffix(".json.tmp")
    with temp.open("w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    os.replace(temp, MANIFEST_FILE)


def scan_files() -> dict[str, str]:
    current = {}
    for path in sorted(NOTES_DIR.rglob("*")):
        if (path.is_file() and path.suffix.lower() in SUPPORTED
                and not path.name.startswith("~$")
                and not path.name.startswith(".")):
            current[str(path.resolve())] = file_hash(path)
    return current


# -------------------- DOCLING --------------------
@lru_cache(maxsize=1)
def get_converter():
    # Lazy-load Docling only when a PDF or DOCX actually needs processing.
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    print("  Loading Docling (first PDF/DOCX of this update)...", flush=True)
    options = PdfPipelineOptions()
    options.do_ocr = True
    options.ocr_options.force_full_page_ocr = False
    options.generate_page_images = False
    options.generate_picture_images = True
    options.images_scale = 1.5
    return DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
    )


def convert_document(path: Path):
    result = get_converter().convert(str(path))
    return result.document


# -------------------- GRANITE --------------------
def describe_image(image: Image.Image | bytes, location: str, nearby_text: str = "") -> str:
    """Call local Ollama Granite for a single visual."""
    global vision_used
    if isinstance(image, bytes):
        image = Image.open(BytesIO(image))
    image = image.convert("RGB")
    image.thumbnail((1600, 1600))
    out = BytesIO()
    image.save(out, format="PNG")
    vision_used = True
    answer = ollama.chat(
        model=VISION_MODEL,
        messages=[{
            "role": "user",
            "content": (
                f"Visual from {location}. Nearby extracted text: {nearby_text[:1500]}\n"
                "Describe ONLY meaningful content visible in this image. "
                "For equations, transcribe symbols, fractions, subscripts and exponents "
                "as accurately as possible. For charts, name axes, units, legends, "
                "visible values and trends. For diagrams, describe relationships. "
                "If a detail is illegible or uncertain, say so. "
                "Ignore purely decorative images."
            ),
            "images": [out.getvalue()],
        }],
        options={"temperature": 0},
        keep_alive="10m",
    )
    return (answer.message.content or "").strip()


def unload_vision() -> None:
    global vision_used
    if not vision_used:
        return
    try:
        ollama.generate(model=VISION_MODEL, prompt="", keep_alive=0)
        print("Granite unloaded from memory.")
    except Exception as exc:
        print(f"Warning: could not unload Granite: {exc}")
    finally:
        vision_used = False


def docling_figures(doc, source: str, filename: str, max_figures: int):
    """Extract PDF / DOCX visual regions where Docling supplies image crops."""
    output = []
    for number, picture in enumerate(doc.pictures[:max_figures], 1):
        image = picture.get_image(doc)
        if image is None:
            continue
        prov = picture.prov or []
        page = getattr(prov[0], "page_no", None) if prov else None
        location = f"{filename}, figure {number}"
        if page is not None:
            location += f", page {page}"
        print(f"  Granite: {location}")
        description = describe_image(image, location)
        if description:
            metadata = {"source": source, "type": "visual", "figure": number}
            if page is not None:
                metadata["page"] = int(page)
            output.append(Document(page_content=description, metadata=metadata))
    return output


# -------------------- POWERPOINT VISUALS --------------------
def iter_slide_shapes(shapes):
    """Visit group contents as well as ordinary slide shapes."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_slide_shapes(shape.shapes)


def slide_needs_vision(slide) -> bool:
    """Detect visual objects without treating text/table-only groups as visuals."""
    visual_types = {
        MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE,
        MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.DIAGRAM,
        MSO_SHAPE_TYPE.IGX_GRAPHIC, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
        MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM,
        MSO_SHAPE_TYPE.LINE,
    }
    for shape in iter_slide_shapes(slide.shapes):
        if getattr(shape, "has_chart", False) or shape.shape_type in visual_types:
            return True
        # SmartArt can be exposed as an unclassified graphic frame by python-pptx.
        # Inspect its DrawingML payload rather than counting table frames as visuals.
        for element in shape.element.iter():
            if element.tag.rsplit("}", 1)[-1] in {"relIds", "oleObj"}:
                return True
    return False


def render_slides_to_images(path: Path) -> list[Image.Image]:
    """Optional whole-slide rendering (requires LibreOffice and pymupdf)."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not Path(soffice).exists() and not shutil.which("soffice"):
        return []
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return []

    with tempfile.TemporaryDirectory(prefix="rag_slides_") as tmp:
        tmpdir = Path(tmp)
        command = [
            soffice, "-env:UserInstallation=file://" + str(tmpdir / "profile"),
            "--headless", "--convert-to", "pdf", "--outdir", str(tmpdir), str(path)
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=180)
        pdf = tmpdir / (path.stem + ".pdf")
        if completed.returncode != 0 or not pdf.exists():
            print("  Slide rendering unavailable; falling back to embedded pictures.")
            return []
        images = []
        with fitz.open(pdf) as pdf_doc:
            for page in pdf_doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                images.append(Image.open(BytesIO(pix.tobytes("png"))).copy())
        return images


def native_chart_text(slide) -> list[str]:
    """Extract native PowerPoint chart values even without slide rendering."""
    items = []
    for shape in iter_slide_shapes(slide.shapes):
        if not getattr(shape, "has_chart", False):
            continue
        try:
            chart = shape.chart
            labels = []
            if chart.plots and chart.plots[0].categories:
                labels = [str(c.label) for c in chart.plots[0].categories]
            for series in chart.series:
                values = [str(v) for v in series.values]
                pairs = ", ".join(
                    f"{labels[i]}: {v}" if i < len(labels) else v
                    for i, v in enumerate(values)
                )
                items.append(f"Chart series {series.name}: {pairs}")
        except Exception as exc:
            print(f"  Could not extract native chart data: {exc}")
    return items


def pptx_content(path: Path, source: str) -> list[Document]:
    """Preserve slide-level citations and send slide visuals to Granite."""
    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    needs_vision = [slide_needs_vision(slide) for slide in slides]
    images = render_slides_to_images(path) if RENDER_FULL_SLIDES and any(needs_vision) else []
    if images:
        print("  Full-slide rendering enabled (including native charts/equations).")
    else:
        print("  Embedded-picture fallback (native drawings may not be visible to Granite).")
    output = []
    visuals = 0
    for slide_num, slide in enumerate(slides, 1):
        text_parts = []
        for shape in iter_slide_shapes(slide.shapes):
            if shape.has_text_frame and shape.text.strip():
                text_parts.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    text_parts.append(" | ".join(cell.text for cell in row.cells))
        text_parts += native_chart_text(slide)
        nearby = "\n".join(text_parts)
        if nearby.strip():
            output.append(Document(
                page_content=nearby,
                metadata={"source": source, "slide": slide_num, "type": "slide_text"}
            ))
        if not needs_vision[slide_num - 1]:
            print(f"  Skipping Granite: {path.name}, slide {slide_num} (text/table only)")
            continue
        if images and slide_num <= len(images):
            if visuals >= MAX_VISUALS_PER_FILE:
                print("  Visual limit reached; remaining visuals not analyzed.")
                continue
            # Full slides can contain more than pictures (charts, equation objects).
            print(f"  Granite: {path.name}, slide {slide_num}")
            description = describe_image(images[slide_num - 1],
                                         f"{path.name}, slide {slide_num}", nearby)
            visuals += 1
            if description:
                output.append(Document(
                    page_content=description,
                    metadata={"source": source, "slide": slide_num, "type": "visual"}
                ))
        elif not images:
            for shape in iter_slide_shapes(slide.shapes):
                if visuals >= MAX_VISUALS_PER_FILE:
                    break
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                print(f"  Granite: {path.name}, slide {slide_num}, embedded picture")
                description = describe_image(shape.image.blob,
                                             f"{path.name}, slide {slide_num}", nearby)
                visuals += 1
                if description:
                    output.append(Document(
                        page_content=description,
                        metadata={"source": source, "slide": slide_num, "type": "visual"}
                    ))
    return output


# -------------------- DOCUMENT EXTRACTION --------------------
def load_single_document(path: Path) -> list[Document]:
    source = str(path.resolve())
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        content = path.read_text(encoding="utf-8")
        return [Document(page_content=content, metadata={"source": source})] if content.strip() else []

    if suffix == ".pptx":
        # Docling parses PPTX structure, but python-pptx preserves slide citations.
        # Extract slide-level content with python-pptx to avoid duplicate embeddings.
        return pptx_content(path, source)

    if suffix in {".pdf", ".docx"}:
        print(f"  Docling: {path.name}")
        doc = convert_document(path)
        markdown = doc.export_to_markdown()
        output = []
        if markdown.strip():
            output.append(Document(
                page_content=markdown,
                metadata={"source": source, "type": "document_text"}
            ))
        output += docling_figures(doc, source, path.name, MAX_VISUALS_PER_FILE)
        return output
    return []


# -------------------- INDEX UPDATES --------------------
def file_chunk_ids(db: Chroma, source: str) -> list[str]:
    return db.get(where={"source": source}, include=[])["ids"]


def index_file(db: Chroma, path: Path, expected_hash: str) -> int:
    """Stage replacement chunks before removing old chunks; roll back on errors."""
    source = str(path.resolve())
    docs = load_single_document(path)
    chunks = splitter.split_documents(docs)
    if not chunks:
        raise ValueError("No extractable content found (or no readable images).")
    if file_hash(path) != expected_hash:
        raise RuntimeError("File changed during processing; retry the update.")
    new_ids = [hashlib.sha256(
        f"{PIPELINE_VERSION}|{source}|{expected_hash}|{i}".encode("utf-8")
    ).hexdigest() for i in range(len(chunks))]
    old_ids = file_chunk_ids(db, source)
    # Clear any leftover partial attempt with the same stable IDs.
    existing_new = set(db.get(ids=new_ids, include=[])["ids"])
    if existing_new:
        db.delete(ids=list(existing_new))
    added = []
    try:
        print(f"  Embedding {len(chunks)} chunks (batches of {BATCH_SIZE})...")
        for i in range(0, len(chunks), BATCH_SIZE):
            batch_ids = new_ids[i:i + BATCH_SIZE]
            db.add_documents(documents=chunks[i:i + BATCH_SIZE], ids=batch_ids)
            added.extend(batch_ids)
            print(f"  Indexed {min(i + BATCH_SIZE, len(chunks))}/{len(chunks)}")
        if file_hash(path) != expected_hash:
            raise RuntimeError("File changed during indexing; retaining previous version.")
        stale_ids = [ident for ident in old_ids if ident not in new_ids]
        if stale_ids:
            db.delete(ids=stale_ids)
        return len(chunks)
    except Exception:
        # Existing previous-version chunks were not deleted yet.
        if added:
            db.delete(ids=added)
        raise


def update_database(db: Chroma) -> None:
    manifest = load_manifest()
    previous = manifest["files"]
    current = scan_files()
    new = [s for s in current if s not in previous]
    changed = [s for s in current if s in previous and current[s] != previous[s]]
    deleted = [s for s in previous if s not in current]
    print(f"\nChanges: {len(new)} new, {len(changed)} modified, {len(deleted)} deleted")
    if not (new or changed or deleted):
        print("Database already up to date.")
        return
    for source in new + changed:
        path = Path(source)
        print(f"\nProcessing: {path.name}")
        try:
            count = index_file(db, path, current[source])
            previous[source] = current[source]
            save_manifest(manifest)
            print(f"  Complete: {count} chunks")
        except Exception as exc:
            print(f"  ERROR: {exc}")
            print("  This file will be retried on the next update.")
    for source in deleted:
        print(f"\nRemoving deleted file: {Path(source).name}")
        try:
            ids = file_chunk_ids(db, source)
            if ids:
                db.delete(ids=ids)
            previous.pop(source, None)
            save_manifest(manifest)
            print(f"  Removed {len(ids)} chunks")
        except Exception as exc:
            print(f"  ERROR: {exc}; will retry on next update.")
    print("\nDatabase update complete.")


def update_and_unload(db: Chroma) -> None:
    try:
        update_database(db)
    finally:
        # Release Docling objects and unload Ollama's vision model.
        get_converter.cache_clear()
        gc.collect()
        unload_vision()


# -------------------- CHAT --------------------
def choose_llm_model() -> str:
    """Choose an installed chat model without downloading or loading one."""
    def field(value, name, default=None):
        return value.get(name, default) if isinstance(value, dict) else getattr(value, name, default)

    def base_name(name):
        return name.rsplit("/", 1)[-1].split(":", 1)[0]

    try:
        response = ollama.list()
        names = sorted({
            name for model in field(response, "models", [])
            if isinstance(name := field(model, "model") or field(model, "name"), str) and name
        })
        excluded = {base_name(EMBED_MODEL), base_name(VISION_MODEL)}
        chat_models = []
        for name in names:
            if base_name(name) in excluded:
                continue
            try:
                capabilities = field(ollama.show(name), "capabilities")
            except Exception:
                capabilities = None  # Older Ollama servers may not expose capabilities.
            if capabilities is not None and "completion" not in capabilities:
                continue
            chat_models.append(name)
    except Exception as exc:
        print(f"\nCould not read Ollama model list: {exc}")
        print(f"Using default model: {LLM_MODEL}")
        return LLM_MODEL

    if not chat_models:
        print(f"\nNo selectable chat models found. Using default: {LLM_MODEL}")
        return LLM_MODEL

    default_model = LLM_MODEL if LLM_MODEL in chat_models else chat_models[0]
    print("\nAVAILABLE CHAT MODELS")
    for number, model in enumerate(chat_models, 1):
        marker = "  [default]" if model == default_model else ""
        print(f"{number}. {model}{marker}")
    while True:
        try:
            choice = input(f"Choose model number/name [Enter = {default_model}]: ").strip()
        except (EOFError, KeyboardInterrupt):
            print(f"\nUsing default model: {default_model}")
            return default_model
        if not choice:
            return default_model
        if choice in chat_models:
            return choice
        if choice.isdecimal() and len(choice) < 10:
            index = int(choice) - 1
            if 0 <= index < len(chat_models):
                return chat_models[index]
        print("Invalid selection. Enter a listed number/name, or press Enter for the default.")


def main() -> None:
    NOTES_DIR.mkdir(parents=True, exist_ok=True)
    db_exists = DB_DIR.exists() and any(DB_DIR.iterdir())
    if db_exists and not MANIFEST_FILE.exists():
        raise RuntimeError(
            "Existing chroma_db has no indexed_files.json. "
            "For a clean start, back up/delete chroma_db and indexed_files.json."
        )
    if not db_exists and MANIFEST_FILE.exists():
        print("Database missing: removing stale indexing history.")
        MANIFEST_FILE.unlink()
    if db_exists:
        load_manifest()  # Verify the pipeline version before opening the DB.

    embeddings = OllamaEmbeddings(model=EMBED_MODEL)
    db = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embeddings,
        persist_directory=str(DB_DIR),
    )
    if not db_exists:
        print("\nYour database is empty.")
        choice = input("Index your documents now? (y/n): ").strip().lower()
    else:
        print("\nExisting database loaded.")
        choice = input(
            "Check my_notes for new, modified or deleted files? (y/n): "
        ).strip().lower()

    if choice in {"y", "yes"}:
        update_and_unload(db)
    else:
        print("Skipping document processing; starting chat with the existing index.")

    selected_model = choose_llm_model()
    llm = ChatOllama(model=selected_model, temperature=0, num_ctx=8192)
    print(f"\nJD_RAG 0.2.0 ready. Indexed chunks: {db._collection.count()}")
    print(f"Chat model: {selected_model}")
    print(f"Embedding model: {EMBED_MODEL}")
    prompt = ChatPromptTemplate.from_messages([
        ("system",
         "Answer using ONLY the retrieved document context. If the information is "
         "missing or uncertain, say so. Treat visual descriptions as imperfect; "
         "do not invent equations or exact chart values. Reference filenames and "
         "page/slide numbers where present.\n\nCONTEXT:\n{context}"),
        ("human", "{question}"),
    ])
    retriever = db.as_retriever(search_kwargs={"k": RETRIEVAL_K})
    print("Ask about your notes. Type 'quit' to exit, or '/update' to rescan.")
    while True:
        try:
            question = input("\nYou: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nGoodbye!")
            break
        if question.lower() in {"quit", "exit"}:
            break
        if question.lower() == "/update":
            update_and_unload(db)
            continue
        if not question:
            continue
        try:
            docs = retriever.invoke(question)
            if not docs:
                print("\nNo indexed documents found. Add files and enter /update.")
                continue
            context_parts = []
            for doc in docs:
                m = doc.metadata
                loc = f"page {m['page']}" if "page" in m else (
                    f"slide {m['slide']}" if "slide" in m else "")
                context_parts.append(
                    f"Source: {Path(m['source']).name} {loc}\n"
                    f"Content type: {m.get('type', 'text')}\n{doc.page_content}"
                )
            message = prompt.invoke({
                "context": "\n\n---\n\n".join(context_parts), "question": question
            })
            response = llm.invoke(message)
            print(f"\n{selected_model}: {response.content}")
            print("\nRetrieved sources:")
            seen = set()
            for doc in docs:
                m = doc.metadata
                label = Path(m["source"]).name
                if "page" in m:
                    label += f", page {m['page']}"
                elif "slide" in m:
                    label += f", slide {m['slide']}"
                if label not in seen:
                    print(f"  - {label}")
                    seen.add(label)
        except Exception as exc:
            print(f"\nError: {exc}")


if __name__ == "__main__":
    main()
