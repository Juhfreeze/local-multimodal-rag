from __future__ import annotations

import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

import ollama
from PIL import Image
from langchain_core.documents import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .config import Settings
from .events import ProgressCallback, ignore_progress


class DocumentProcessor:
    """Document extraction and vision; use one instance per indexing session.

    Operations are synchronous. A GUI should run them on a worker, serially.
    """
    def __init__(self, settings: Settings, report: ProgressCallback = ignore_progress):
        self.settings = settings
        self.report = report
        self.vision_used = False
        self._converter = None

    def clear_cache(self):
        self._converter = None

    def get_converter(self):
        if self._converter is not None:
            return self._converter
        # Lazy-load Docling only when a PDF or DOCX actually needs processing.
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions
        from docling.document_converter import DocumentConverter, PdfFormatOption

        self.report("  Loading Docling (first PDF/DOCX of this update)...")
        options = PdfPipelineOptions()
        options.do_ocr = True
        options.ocr_options.force_full_page_ocr = False
        options.generate_page_images = False
        options.generate_picture_images = True
        options.images_scale = 1.5
        self._converter = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
        )
        return self._converter

    def convert_document(self, path: Path):
        result = self.get_converter().convert(str(path))
        return result.document

    def describe_image(self, image: Image.Image | bytes, location: str, nearby_text: str = "") -> str:
        """Call local Ollama Granite for a single visual."""
        if isinstance(image, bytes):
            image = Image.open(BytesIO(image))
        image = image.convert("RGB")
        image.thumbnail((1600, 1600))
        out = BytesIO()
        image.save(out, format="PNG")
        self.vision_used = True
        answer = ollama.chat(
            model=self.settings.vision_model,
            messages=[{
                "role": "user",
                "content": (
                    f"Visual from {location}. Nearby extracted text: {nearby_text[:1500]}\n"
                    "Describe ONLY meaningful content visible in this image. "
                    "For equations, transcribe symbols, fractions, subscripts and exponents "
                    "as accurately as possible. For charts, name axes, units, legends, "
                    "visible values and trends. For diagrams, describe relationships. "
                    "If a detail is illegible or uncertain, say so. "
                    "Ignore purely decorative images."
                ),
                "images": [out.getvalue()],
            }],
            options={"temperature": 0},
            keep_alive="10m",
        )
        return (answer.message.content or "").strip()

    def unload_vision(self) -> None:
        if not self.vision_used:
            return
        try:
            ollama.generate(model=self.settings.vision_model, prompt="", keep_alive=0)
            self.report("Granite unloaded from memory.")
        except Exception as exc:
            self.report(f"Warning: could not unload Granite: {exc}")
        finally:
            self.vision_used = False

    def docling_figures(self, doc, source: str, filename: str, max_figures: int):
        """Extract PDF / DOCX visual regions where Docling supplies image crops."""
        output = []
        for number, picture in enumerate(doc.pictures[:max_figures], 1):
            image = picture.get_image(doc)
            if image is None:
                continue
            prov = picture.prov or []
            page = getattr(prov[0], "page_no", None) if prov else None
            location = f"{filename}, figure {number}"
            if page is not None:
                location += f", page {page}"
            self.report(f"  Granite: {location}")
            description = self.describe_image(image, location)
            if description:
                metadata = {"source": source, "type": "visual", "figure": number}
                if page is not None:
                    metadata["page"] = int(page)
                output.append(Document(page_content=description, metadata=metadata))
        return output

    def iter_slide_shapes(self, shapes):
        """Visit group contents as well as ordinary slide shapes."""
        for shape in shapes:
            yield shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self.iter_slide_shapes(shape.shapes)

    def slide_needs_vision(self, slide) -> bool:
        """Detect visual objects without treating text/table-only groups as visuals."""
        visual_types = {
            MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE,
            MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.DIAGRAM,
            MSO_SHAPE_TYPE.IGX_GRAPHIC, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
            MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
            MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM,
            MSO_SHAPE_TYPE.LINE,
        }
        for shape in self.iter_slide_shapes(slide.shapes):
            if getattr(shape, "has_chart", False) or shape.shape_type in visual_types:
                return True
            # SmartArt can be exposed as an unclassified graphic frame by python-pptx.
            # Inspect its DrawingML payload rather than counting table frames as visuals.
            for element in shape.element.iter():
                if element.tag.rsplit("}", 1)[-1] in {"pic", "relIds", "oleObj"}:
                    return True
        return False

    def render_slides_to_images(self, path: Path) -> list[Image.Image]:
        """Optional whole-slide rendering (requires LibreOffice and pymupdf)."""
        soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        if not Path(soffice).exists() and not shutil.which("soffice"):
            return []
        try:
            import fitz  # PyMuPDF
        except ImportError:
            return []

        with tempfile.TemporaryDirectory(prefix="rag_slides_") as tmp:
            tmpdir = Path(tmp)
            command = [
                soffice, "-env:UserInstallation=file://" + str(tmpdir / "profile"),
                "--headless", "--convert-to", "pdf", "--outdir", str(tmpdir), str(path)
            ]
            completed = subprocess.run(command, capture_output=True, text=True, timeout=180)
            pdf = tmpdir / (path.stem + ".pdf")
            if completed.returncode != 0 or not pdf.exists():
                self.report("  Slide rendering unavailable; falling back to embedded pictures.")
                return []
            images = []
            with fitz.open(pdf) as pdf_doc:
                for page in pdf_doc:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    images.append(Image.open(BytesIO(pix.tobytes("png"))).copy())
            return images

    def native_chart_text(self, slide) -> list[str]:
        """Extract native PowerPoint chart values even without slide rendering."""
        items = []
        for shape in self.iter_slide_shapes(slide.shapes):
            if not getattr(shape, "has_chart", False):
                continue
            try:
                chart = shape.chart
                labels = []
                if chart.plots and chart.plots[0].categories:
                    labels = [str(c.label) for c in chart.plots[0].categories]
                for series in chart.series:
                    values = [str(v) for v in series.values]
                    pairs = ", ".join(
                        f"{labels[i]}: {v}" if i < len(labels) else v
                        for i, v in enumerate(values)
                    )
                    items.append(f"Chart series {series.name}: {pairs}")
            except Exception as exc:
                self.report(f"  Could not extract native chart data: {exc}")
        return items

    def pptx_content(self, path: Path, source: str) -> list[Document]:
        """Preserve slide-level citations and send slide visuals to Granite."""
        presentation = Presentation(str(path))
        slides = list(presentation.slides)
        needs_vision = [self.slide_needs_vision(slide) for slide in slides]
        images = self.render_slides_to_images(path) if self.settings.render_full_slides and any(needs_vision) else []
        if images:
            self.report("  Full-slide rendering enabled (including native charts/equations).")
        elif any(needs_vision):
            self.report("  Embedded-picture fallback (native drawings may not be visible to Granite).")
        else:
            self.report("  No slide visuals detected; skipping slide rendering and Granite.")
        output = []
        visuals = 0
        for slide_num, slide in enumerate(slides, 1):
            text_parts = []
            for shape in self.iter_slide_shapes(slide.shapes):
                if shape.has_text_frame and shape.text.strip():
                    text_parts.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        text_parts.append(" | ".join(cell.text for cell in row.cells))
            text_parts += self.native_chart_text(slide)
            nearby = "\n".join(text_parts)
            if nearby.strip():
                output.append(Document(
                    page_content=nearby,
                    metadata={"source": source, "slide": slide_num, "type": "slide_text"}
                ))
            if not needs_vision[slide_num - 1]:
                self.report(f"  Skipping Granite: {path.name}, slide {slide_num} (text/table only)")
                continue
            if images and slide_num <= len(images):
                if visuals >= self.settings.max_visuals_per_file:
                    self.report("  Visual limit reached; remaining visuals not analyzed.")
                    continue
                # Full slides can contain more than pictures (charts, equation objects).
                self.report(f"  Granite: {path.name}, slide {slide_num}")
                description = self.describe_image(images[slide_num - 1],
                                             f"{path.name}, slide {slide_num}", nearby)
                visuals += 1
                if description:
                    output.append(Document(
                        page_content=description,
                        metadata={"source": source, "slide": slide_num, "type": "visual"}
                    ))
            elif not images:
                for shape in self.iter_slide_shapes(slide.shapes):
                    if visuals >= self.settings.max_visuals_per_file:
                        break
                    if not hasattr(shape, "image"):
                        continue
                    self.report(f"  Granite: {path.name}, slide {slide_num}, embedded picture")
                    description = self.describe_image(shape.image.blob,
                                                 f"{path.name}, slide {slide_num}", nearby)
                    visuals += 1
                    if description:
                        output.append(Document(
                            page_content=description,
                            metadata={"source": source, "slide": slide_num, "type": "visual"}
                        ))
        return output

    def load_single_document(self, path: Path) -> list[Document]:
        source = str(path.resolve())
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            content = path.read_text(encoding="utf-8")
            return [Document(page_content=content, metadata={"source": source})] if content.strip() else []

        if suffix == ".pptx":
            # Docling parses PPTX structure, but python-pptx preserves slide citations.
            # Extract slide-level content with python-pptx to avoid duplicate embeddings.
            return self.pptx_content(path, source)

        if suffix in {".pdf", ".docx"}:
            self.report(f"  Docling: {path.name}")
            doc = self.convert_document(path)
            markdown = doc.export_to_markdown()
            output = []
            if markdown.strip():
                output.append(Document(
                    page_content=markdown,
                    metadata={"source": source, "type": "document_text"}
                ))
            output += self.docling_figures(doc, source, path.name, self.settings.max_visuals_per_file)
            return output
        return []
